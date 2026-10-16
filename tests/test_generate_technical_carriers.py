"""CP-30D: editable-neutral print grids and genuine OOXML slide objects.

Actual customer acceptance and native editability are NOT certified here.
"""
import json
import tempfile
import unittest
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pypdf import PdfReader

from scripts.generate_technical_carriers import (
    OUTPUT, ROOT, SLIDE_TYPES, CARRIERS, PAPER_MM,
    specs, svg_carrier, generate, verify, px_300,
)


class TechnicalCarrierTests(unittest.TestCase):
    def test_13_independent_tz_carriers_and_11_posters(self):
        docs=specs()
        self.assertEqual(len(docs),13)
        self.assertEqual(sum(x.id.startswith("POSTER-") for x in docs),11)
        self.assertEqual(sum(x.id.startswith("PRICE-HOURS-") for x in docs),2)
        self.assertEqual(len({x.id for x in docs}),13)
        matrix=(ROOT/"docs/tz-audit/requirements.psv").read_text()
        for x in docs:
            self.assertIn(x.tz,matrix)
            self.assertEqual(x.trim_mm,PAPER_MM[x.paper])
            self.assertEqual(x.artboard_mm,(x.trim_mm[0]+6,x.trim_mm[1]+6))
            self.assertEqual(x.min_raster_pixels_bleed,
                             (px_300(x.artboard_mm[0]),px_300(x.artboard_mm[1])))
            self.assertGreater(x.min_raster_pixels_trim[0],2000)
        with self.assertRaisesRegex(ValueError,"Unsafe"):
            specs(bleed=30)
        with self.assertRaisesRegex(ValueError,"Unsafe"):
            specs(safe=0)

    def test_svg_is_editable_and_has_guides_separate_from_content(self):
        first=specs()[0]
        root=ET.fromstring(svg_carrier(first))
        ns="{http://www.w3.org/2000/svg}"
        mb="{https://musbread.local/technical-draft}"
        self.assertEqual(root.attrib[mb+"status"],"TECHNICAL_DRAFT_ONLY")
        self.assertEqual(root.attrib["width"],"303mm")
        ids={n.get("id") for n in root.iter()}
        for target in ("trim-guide","safe-guide","editable-content","type-headline",
                       "draft-watermark","image-placeholder","crop-marks"):
            self.assertIn(target,ids)
        texts=[n.text for n in root.iter() if n.tag==ns+"text"]
        self.assertTrue(any("НЕ УТВЕРЖДЕНО" in (t or "") for t in texts))
        self.assertGreaterEqual(len(texts),6)
        self.assertNotIn("<image",svg_carrier(first))
        self.assertNotIn('href="http',svg_carrier(first))
        self.assertNotIn("<script",svg_carrier(first))

    def test_pptx_both_ratios_have_real_independent_text_and_shapes(self):
        result=verify(OUTPUT)
        self.assertEqual(result["neutral_svg_print_layouts"],13)
        self.assertEqual(result["pptx_editable_slide_layouts"],16)
        for ratio in ("4x3","16x9"):
            deck=Presentation(str(OUTPUT/f"TECH_SLIDES_{ratio}.pptx"))
            self.assertEqual(len(deck.slides),8)
            for n,(code,_) in enumerate(SLIDE_TYPES):
                slide=deck.slides[n]
                labels=[x.text for x in slide.shapes if x.has_text_frame]
                self.assertTrue(any("ТЕХНИЧЕСКАЯ РАСКЛАДКА" in t for t in labels))
                self.assertTrue(any("DRAFT" in t for t in labels))
            six=deck.slides[6]
            boxlabels=[s.text for s in six.shapes if s.has_text_frame]
            for n in range(1,7):
                self.assertTrue(any(f"БЛОК {n}" in l for l in boxlabels))

    def test_pdf_previews_are_physical_artboards_including_bleed(self):
        for x in specs():
            filename=OUTPUT/f"TECH_{x.id}.pdf"
            self.assertTrue(filename.is_file(), filename.name)
            page=PdfReader(str(filename)).pages[0]
            self.assertAlmostEqual(float(page.mediabox.width)*25.4/72,
                                   x.artboard_mm[0],places=1)
            self.assertAlmostEqual(float(page.mediabox.height)*25.4/72,
                                   x.artboard_mm[1],places=1)

    def test_generated_manifest_is_explicitly_not_approved(self):
        contents=json.loads((OUTPUT/"TECHNICAL_MANIFEST.json").read_text())
        self.assertEqual(contents["status"],"NOT_APPROVED_NOT_FINAL_DESIGN")
        self.assertEqual(len(contents["paper_layouts"]),13)
        self.assertEqual(len(contents["pptx_decks"]),2)
        self.assertEqual(set(contents["pptx_decks"]),
                         {"TECH_SLIDES_4x3.pptx","TECH_SLIDES_16x9.pptx"})
        for x in contents["paper_layouts"]:
            self.assertEqual(x["status"],"PREAWARD_TECHNICAL_NOT_CUSTOMER_APPROVED")

    def test_modified_and_missing_outputs_fail_closed(self):
        with tempfile.TemporaryDirectory() as folder:
            output=Path(folder)
            generate(output)
            self.assertEqual(verify(output)["neutral_svg_print_layouts"],13)
            file=output/"TECH_POSTER-CHILD-A4.svg"
            file.write_text(file.read_text().replace("НЕ УТВЕРЖДЕНО","УТВЕРЖДЕНО"))
            with self.assertRaisesRegex(ValueError,"modified"):
                verify(output)
            file.unlink()
            with self.assertRaisesRegex(ValueError,"Missing print layout"):
                verify(output)


if __name__=="__main__":
    unittest.main()
