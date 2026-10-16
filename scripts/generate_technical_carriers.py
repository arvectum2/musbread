#!/usr/bin/env python3
"""MusBread CP-30D — neutral, editable pre-interview carrier geometry.

Generated technical drafts are NOT customer-approved design or TZ deliverables.
A2/A3/A4 sizes are TZ-based; 3mm bleed and 10mm safe margin are proposals.
Builds 13 SVG + two fully editable, eight-slide PowerPoint decks; no museum data.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import subprocess
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/design/technical-drafts"
PAPER_MM = {"A2": (420, 594), "A3": (297, 420), "A4": (210, 297)}
BLEED_MM = 3
SAFE_MM = 10

# TZ references are exact identifiers from docs/tz-audit/requirements.psv.
CARRIERS = (
    ("POSTER-TEMP-A3", "TZ-038", "Временная выставка", "A3"),
    ("POSTER-TEMP-A2", "TZ-038", "Временная выставка", "A2"),
    ("POSTER-CHILD-A4", "TZ-039", "Детский мастер-класс", "A4"),
    ("POSTER-CHILD-A3", "TZ-039", "Детский мастер-класс", "A3"),
    ("POSTER-CHILD-A2", "TZ-039", "Детский мастер-класс", "A2"),
    ("POSTER-SCIENCE-A4", "TZ-040", "Научное событие", "A4"),
    ("POSTER-SCIENCE-A3", "TZ-040", "Научное событие", "A3"),
    ("POSTER-SCIENCE-A2", "TZ-040", "Научное событие", "A2"),
    ("POSTER-LEISURE-A4", "TZ-041", "Развлекательное событие", "A4"),
    ("POSTER-LEISURE-A3", "TZ-041", "Развлекательное событие", "A3"),
    ("POSTER-LEISURE-A2", "TZ-041", "Развлекательное событие", "A2"),
    ("PRICE-HOURS-A4", "TZ-042", "Часы работы и цены", "A4"),
    ("PRICE-HOURS-A3", "TZ-042", "Часы работы и цены", "A3"),
)
SLIDE_TYPES = (
    ("TITLE", "Титульный слайд"),
    ("CLOSING", "Завершение"),
    ("TEXT", "Объёмный текст"),
    ("IMAGE", "Крупное изображение"),
    ("MIXED", "Текст и изображение"),
    ("FOUR", "Четыре текстовых блока"),
    ("SIX", "Шесть текстовых блоков"),
    ("DATA", "Таблица и данные"),
)

@dataclass(frozen=True)
class Carrier:
    id: str
    tz: str
    title: str
    paper: str
    trim_mm: tuple[int, int]
    bleed_mm: int
    safe_mm: int
    artboard_mm: tuple[int, int]
    min_raster_pixels_trim: tuple[int, int]
    min_raster_pixels_bleed: tuple[int, int]

def px_300(mm: float) -> int:
    return math.ceil(mm * 300 / 25.4)

def specs(bleed: int=BLEED_MM, safe: int=SAFE_MM) -> list[Carrier]:
    if not 0<=bleed<=10 or not 3<=safe<=25:
        raise ValueError("Unsafe or implausible prepress defaults")
    result=[]
    for id, tz, title, paper in CARRIERS:
        w,h=PAPER_MM[paper]
        a,b=w+2*bleed,h+2*bleed
        result.append(Carrier(
            id,tz,title,paper,(w,h),bleed,safe,(a,b),
            (px_300(w),px_300(h)),(px_300(a),px_300(b))))
    return result

def svg_carrier(carrier:Carrier)->str:
    art_w,art_h=carrier.artboard_mm
    w,h=carrier.trim_mm
    b=carrier.bleed_mm
    safety=carrier.safe_mm
    left=b+safety
    right=b+w-safety
    top=b+safety
    innerw=right-left
    # Use real SVG text and independent shapes. Nothing is flattened.
    heading=escape(carrier.title)
    lines=[
        '<?xml version="1.0" encoding="utf-8"?>',
        '<svg xmlns="http://www.w3.org/2000/svg" xmlns:mb="https://musbread.local/technical-draft" '
        f'width="{art_w}mm" height="{art_h}mm" viewBox="0 0 {art_w} {art_h}" '
        f'mb:paper="{carrier.paper}" mb:tz="{carrier.tz}" '
        f'mb:trim-width-mm="{w}" mb:trim-height-mm="{h}" '
        f'mb:bleed-mm="{b}" mb:safe-mm="{safety}" '
        f'mb:carrier="{carrier.id}" mb:status="TECHNICAL_DRAFT_ONLY">',
        '<title>MusBread — нейтральная техническая сетка, не утверждённый дизайн</title>',
        '<desc>Редактируемая SVG-заготовка. ТЗ задаёт A2/A3/A4; вылеты 3 мм и зона 10 мм — только предложение для согласования с типографией.</desc>',
        '<g id="background-editable"><rect width="100%" height="100%" fill="#ffffff"/></g>',
        f'<g id="guides-trim" fill="none" stroke="#6b7280" stroke-width="0.35">'
        f'<rect id="trim-guide" x="{b}" y="{b}" width="{w}" height="{h}"/></g>',
        f'<g id="guides-safe" fill="none" stroke="#b45309" stroke-width="0.35" stroke-dasharray="3 2">'
        f'<rect id="safe-guide" x="{left}" y="{top}" width="{innerw}" height="{h-2*safety}"/></g>',
    ]
    # Four crop marks outside the trim, clear from bleed edge.
    crop_lines=[]
    for x in (b,b+w):
        crop_lines.append(f'<path d="M{x} 0.6 V{max(0,b-0.6)} M{x} {b+h+0.6} V{art_h-0.6}"/>')
    for y in (b,b+h):
        crop_lines.append(f'<path d="M0.6 {y} H{max(0,b-0.6)} M{b+w+0.6} {y} H{art_w-0.6}"/>')
    lines += ['<g id="crop-marks" stroke="#111827" stroke-width="0.25" fill="none">',
              *crop_lines,'</g>']
    title_y=top+22
    date_y=title_y+30
    box_top=date_y+14
    box_bottom=b+h-safety-31
    imageh=max(18,(box_bottom-box_top)*.55)
    lines += [
        f'<g id="editable-content" font-family="Arial, sans-serif" fill="#111827">',
        f'<text id="type-category" x="{left}" y="{top+6}" font-size="5">РАЗДЕЛ / РУБРИКА — ЗАМЕНИТЬ</text>',
        f'<text id="type-headline" x="{left}" y="{title_y}" font-size="11" font-weight="bold">{heading}</text>',
        f'<text id="type-datetime" x="{left}" y="{date_y}" font-size="5">ДАТА · ВРЕМЯ · МЕСТО — ПРИМЕР ПОЛЕЙ</text>',
        f'<text id="type-body" x="{left}" y="{box_top+imageh+12}" font-size="6">[УТВЕРЖДЁННОЕ ОПИСАНИЕ МЕРОПРИЯТИЯ]</text>',
        f'<text id="type-footer" x="{left}" y="{b+h-safety-13}" font-size="4.5">[КОНТАКТЫ / БИЛЕТЫ / ОФИЦИАЛЬНЫЕ ДАННЫЕ]</text>',
        f'<text id="draft-watermark" x="{left}" y="{b+h-safety-5}" font-size="4" fill="#9a3412">ТЕХНИЧЕСКИЙ ШАБЛОН · НЕ УТВЕРЖДЕНО</text>',
        '</g>',
        f'<g id="editable-placeholders" fill="none" stroke="#94a3b8" stroke-dasharray="3 2" stroke-width=".4">'
        f'<rect id="image-placeholder" x="{left}" y="{box_top}" width="{innerw}" height="{imageh}" rx="1"/>'
        f'</g>',
        f'<text id="image-label" x="{left+5}" y="{box_top+10}" font-family="Arial" font-size="5" fill="#475569">[МЕСТО ДЛЯ ПРАВООЧИЩЕННОГО ИЗОБРАЖЕНИЯ]</text>',
        '</svg>',
    ]
    return "\n".join(lines)+"\n"

def build_slides(path:Path,ratio:str)->None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt
    except ImportError as err:
        raise RuntimeError("Install python-pptx==1.0.2 from scripts/requirements-design-tools.txt") from err
    if ratio not in ("4x3","16x9"):
        raise ValueError(ratio)
    ppt=Presentation()
    ppt.slide_width=Inches(12 if ratio=="4x3" else 13.333333)
    ppt.slide_height=Inches(9 if ratio=="4x3" else 7.5)
    sw=ppt.slide_width / 914400
    sh=ppt.slide_height / 914400
    ink=RGBColor(30,41,59)
    gray=RGBColor(100,116,139)
    border=RGBColor(148,163,184)
    paper=RGBColor(248,250,252)

    def text(slide,x,y,w,h,value,size=19,bold=False,light=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=shape.text_frame
        tf.word_wrap=True
        tf.text=value
        for para in tf.paragraphs:
            for run in para.runs:
                run.font.size=Pt(size)
                run.font.bold=bold
                run.font.color.rgb=gray if light else ink
                run.font.name="Arial"
        return shape

    def shape_box(slide,x,y,w,h,caption,kind="rectangle"):
        obj=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if kind=="rounded" else MSO_SHAPE.RECTANGLE,
                 Inches(x),Inches(y),Inches(w),Inches(h))
        obj.fill.solid()
        obj.fill.fore_color.rgb=paper
        obj.line.color.rgb=border
        obj.line.width=Pt(1)
        text(slide,x+.15,y+.13,max(w-.3,.4),min(h-.2,.7),caption,12,light=True)
        return obj

    def base(code,title,index):
        slide=ppt.slides.add_slide(ppt.slide_layouts[6])
        # All objects are editable PowerPoint shapes and text, not images.
        text(slide,.55,.24,sw-1.1,.28,"MUSBREAD · ТЕХНИЧЕСКАЯ РАСКЛАДКА · НЕ ФИРМЕННЫЙ СТИЛЬ",10,light=True)
        text(slide,.55,.76,sw-1.1,.53,title,26,bold=True)
        rule=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.55),Inches(1.37),Inches(sw-1.1),Inches(.015))
        rule.fill.solid();rule.fill.fore_color.rgb=border;rule.line.fill.background()
        text(slide,.55,sh-.38,sw-1.1,.22,f"TZ-054–058  /  {ratio}  /  {code}  /  {index:02d}  /  DRAFT",9,light=True)
        return slide

    content_y=1.62
    content_h=sh-content_y-.68
    margin=.55
    content_w=sw-2*margin
    gap=.17
    for i,(code,title) in enumerate(SLIDE_TYPES,1):
        slide=base(code,title,i)
        if code in ("TITLE","CLOSING"):
            text(slide,.75,sh*.35,sw-1.5,.72,
                 "[НАЗВАНИЕ ПРЕЗЕНТАЦИИ]" if code=="TITLE" else "[ИТОГ / КОНТАКТЫ]",
                 28,bold=True)
            text(slide,.76,sh*.35+1,sw-1.6,.45,"[ПОДЗАГОЛОВОК И ДАТА]",16,light=True)
        elif code=="TEXT":
            for n in range(3):
                text(slide,.7,content_y+n*content_h/3,content_w-.3,content_h/3-.15,
                     f"{n+1}. [РЕДАКТИРУЕМЫЙ АБЗАЦ · ЗАМЕНИТЬ ТЕКСТ МУЗЕЯ]",17)
        elif code=="IMAGE":
            shape_box(slide,.7,content_y,content_w-.3,content_h-.1,
                      "[ИЗОБРАЖЕНИЕ: ЛИЦЕНЗИЯ И ALT ПЕРЕД ИСПОЛЬЗОВАНИЕМ]")
        elif code=="MIXED":
            col=(content_w-gap)/2
            shape_box(slide,.7,content_y,col,content_h-.1,"[ИЗОБРАЖЕНИЕ]")
            text(slide,.7+col+gap,content_y+.15,col-.15,content_h-.35,
                 "[РЕДАКТИРУЕМЫЙ ТЕКСТ И ПОДПИСЬ]",18)
        elif code in ("FOUR","SIX"):
            count=4 if code=="FOUR" else 6
            cols=2 if code=="FOUR" else 3
            rows=count//cols
            w=(content_w-gap*(cols-1))/cols
            h=(content_h-gap*(rows-1))/rows
            for n in range(count):
                x=.7+(n%cols)*(w+gap)
                y=content_y+(n//cols)*(h+gap)
                shape_box(slide,x,y,w,h,f"[БЛОК {n+1}]")
                text(slide,x+.15,y+.52,w-.3,min(h-.6,.5),"ТЕКСТ / РЕДАКТИРОВАТЬ",11)
        elif code=="DATA":
            cols=3
            w=(content_w-gap*(cols-1))/cols
            for n in range(cols):
                shape_box(slide,.7+n*(w+gap),content_y,w,content_h*.35,f"[МЕТРИКА {n+1}]")
            shape_box(slide,.7,content_y+content_h*.43,content_w,content_h*.52,
                      "[ДИАГРАММА ИЛИ ТАБЛИЦА: ОТДЕЛЬНЫЕ ОБЪЕКТЫ]")
        else:raise AssertionError(code)
    path.parent.mkdir(parents=True,exist_ok=True)
    ppt.save(str(path))

def manifest_for(result:list[dict])->dict:
    return {
        "status":"NOT_APPROVED_NOT_FINAL_DESIGN",
        "source_reference":"Original TZ-038..042, TZ-054..058 for sizes/type counts; bleed and safe are proposed defaults",
        "paper_layouts":result,
        "pptx_decks":[f"TECH_SLIDES_{ratio}.pptx" for ratio in ("4x3","16x9")],
    }


def print_specs_markdown(result:list[dict])->str:
    lines=[
        "# MusBread — technical print geometries, NOT approved art",
        "",
        "A2/A3/A4 trim sizes are required by original TZ-038..042. "
        "**3 mm bleed and 10 mm inner safe zone are preliminary assumptions**, not parameters confirmed by the museum or print shop.",
        "The calculated 300 dpi dimensions apply to a FUTURE **raster at full size**; "
        "the current SVG/PDF vector drafts are not 300 dpi raster print masters and are not CMYK-certified.",
        "",
        "| Carrier ID | TZ | Paper | Trim, mm | With 3 mm bleed, mm | Safe width × height, mm | Raster trim at 300 dpi, px | Raster with bleed at 300 dpi, px |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for row in result:
        tw,th=row["trim_mm"]
        aw,ah=row["artboard_mm"]
        pw,ph=row["min_raster_pixels_trim"]
        bw,bh=row["min_raster_pixels_bleed"]
        lines.append(f'| {row["id"]} | {row["tz"]} | {row["paper"]} | {tw} × {th} | '
                     f'{aw} × {ah} | {tw-2*row["safe_mm"]} × {th-2*row["safe_mm"]} | '
                     f'{pw} × {ph} | {bw} × {bh} |')
    lines += [
        "",
        "## Before any production print run",
        "",
        "1. Confirm bleed and inside safe zone separately for each item and print house.",
        "2. Confirm CMYK profile, coated/uncoated stock, black generation, Pantone if used, color proof and media.",
        "3. Replace ALL bracket placeholders with museum-approved copy, images and actual license/rights.",
        "4. Confirm actual physical dimensions, printable safe zone, trim/crop marks, typography and QR readability.",
        "5. Open editable SVG in agreed vector editor and PDF preview on another device.",
        "6. Re-export actual artwork in formats requested by TZ and verify a raster is >=300 dpi at output size.",
        "7. Receive written customer approval and archive versions and transfer evidence.",
        "",
        "**No official design files, identity system, print sign-off or exclusive-rights transfer has been generated by these technical drafts.**",
        "",
    ]
    return "\n".join(lines)


def generate(output:Path,bleed:int=BLEED_MM,safe:int=SAFE_MM,with_pdf:bool=False)->list[dict]:
    output.mkdir(parents=True,exist_ok=True)
    result=[]
    for carrier in specs(bleed,safe):
        source=output/f"TECH_{carrier.id}.svg"
        source.write_text(svg_carrier(carrier),encoding="utf-8")
        if with_pdf:
            pdf=source.with_suffix(".pdf")
            subprocess.run(["rsvg-convert","--format","pdf","--output",str(pdf),str(source)],
                           check=True,capture_output=True,timeout=12)
        result.append({**asdict(carrier),"source":source.name,
                       "status":"PREAWARD_TECHNICAL_NOT_CUSTOMER_APPROVED"})
    for ratio in ("4x3","16x9"):
        build_slides(output/f"TECH_SLIDES_{ratio}.pptx",ratio)
    m=manifest_for(result)
    (output/"TECHNICAL_MANIFEST.json").write_text(
        json.dumps(m,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    (output/"PRINT_SPECS.md").write_text(print_specs_markdown(result),encoding="utf-8")
    return result

def verify(output:Path,bleed:int=BLEED_MM,safe:int=SAFE_MM)->dict:
    from xml.etree import ElementTree
    from pptx import Presentation
    actual=specs(bleed,safe)
    svg_n=0
    expected_result=[
        {**asdict(c),"source":f"TECH_{c.id}.svg",
         "status":"PREAWARD_TECHNICAL_NOT_CUSTOMER_APPROVED"}
        for c in actual
    ]
    if (output/"TECHNICAL_MANIFEST.json").read_text(encoding="utf-8") != (
            json.dumps(manifest_for(expected_result),ensure_ascii=False,indent=2)+"\n"):
        raise ValueError("Technical manifest is missing/stale/changed or claims a different status")
    if (output/"PRINT_SPECS.md").read_text(encoding="utf-8") != print_specs_markdown(expected_result):
        raise ValueError("300 dpi print specifications report is missing or stale")
    for carrier in actual:
        path=output/f"TECH_{carrier.id}.svg"
        if not path.is_file():
            raise ValueError(f"Missing print layout {path}")
        root=ElementTree.parse(path).getroot()
        if root.attrib.get("width")!=f"{carrier.artboard_mm[0]}mm":
            raise ValueError(f"Wrong full width in {path.name}")
        if root.attrib.get("height")!=f"{carrier.artboard_mm[1]}mm":
            raise ValueError(f"Wrong full height in {path.name}")
        if root.attrib.get("{https://musbread.local/technical-draft}status")!="TECHNICAL_DRAFT_ONLY":
            raise ValueError("SVG missing explicit not-approved status")
        ids={x.attrib.get("id") for x in root.iter()}
        if not {"trim-guide","safe-guide","editable-content","image-placeholder","draft-watermark"}.issubset(ids):
            raise ValueError("SVG lost independent editable guides/content")
        if path.read_text(encoding="utf-8")!=svg_carrier(carrier):
            raise ValueError("Technical SVG was modified without regenerating through source")
        pdf=path.with_suffix(".pdf")
        if pdf.is_file():
            from pypdf import PdfReader
            try:
                pages=PdfReader(str(pdf)).pages
                if len(pages)!=1:
                    raise ValueError("Technical print PDF must contain exactly one page")
                box=pages[0].mediabox
                pw=float(box.width)*25.4/72
                ph=float(box.height)*25.4/72
                if abs(pw-carrier.artboard_mm[0])>.15 or abs(ph-carrier.artboard_mm[1])>.15:
                    raise ValueError("PDF page geometry does not match bleed-inclusive SVG dimensions")
            except Exception as err:
                raise ValueError("PDF technical preview invalid: "+str(pdf)) from err
        svg_n+=1
    presentations={}
    for ratio in ("4x3","16x9"):
        p=Presentation(str(output/f"TECH_SLIDES_{ratio}.pptx"))
        if len(p.slides)!=8:
            raise ValueError("Need 8 editable slide types in each aspect")
        aspect=p.slide_width / p.slide_height
        expected=(4/3) if ratio=="4x3" else (16/9)
        if abs(aspect-expected)>0.001:
            raise ValueError(f"Wrong PPTX ratio {ratio}")
        counts={}
        for (code,_),slide in zip(SLIDE_TYPES,p.slides):
            textboxes=[x for x in slide.shapes if x.has_text_frame and x.text.strip()]
            if len(textboxes)<4:
                raise ValueError("Text is not preserved as editable text shapes")
            boxes=[x for x in slide.shapes if x.shape_type==1]
            if code=="FOUR" and len(boxes)<5:
                raise ValueError("Missing 4 block shapes")
            if code=="SIX" and len(boxes)<7:
                raise ValueError("Missing 6 block shapes")
            counts[code]={"independent_text_shapes":len(textboxes),"vector_shapes":len(boxes)}
        presentations[ratio]=counts
    return {"neutral_svg_print_layouts":svg_n,"pptx_decks":2,
            "pptx_editable_slide_layouts":16,"print_bleed_mm":bleed,
            "print_safe_margin_mm":safe,
            "status":"ALL_TECHNICAL_ONLY_NOT_MUSEUM_APPROVED",
            "presentations":presentations}

def main()->int:
    parser=argparse.ArgumentParser()
    parser.add_argument("--output",type=Path,default=OUTPUT)
    parser.add_argument("--bleed-mm",type=int,default=BLEED_MM)
    parser.add_argument("--safe-mm",type=int,default=SAFE_MM)
    parser.add_argument("--with-pdf-previews",action="store_true")
    parser.add_argument("--verify",action="store_true")
    args=parser.parse_args()
    if not args.verify:
        generate(args.output,args.bleed_mm,args.safe_mm,args.with_pdf_previews)
    print("TECHNICAL CARRIER VERIFICATION:",json.dumps(
        verify(args.output,args.bleed_mm,args.safe_mm),ensure_ascii=False))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
